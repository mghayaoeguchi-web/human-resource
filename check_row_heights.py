#!/usr/bin/env python3
"""
Excelシートの各行について、セル内テキスト量（文字数・フォントサイズ・列幅・折り返し）
から必要な行の高さを見積もり、実際に設定されている行の高さと比較する。
不足している行を検出し、--fix指定時は自動で高さを補正する。

想定：日本語（全角文字中心）のMeiryoフォント、wrap_text=Trueのセルが対象。
"""
import sys
import argparse
import openpyxl
from openpyxl.utils import get_column_letter


def merged_width_units(ws, cell):
    """セルが属する結合範囲の合計列幅（Excel文字単位）を返す。結合されていなければ単一列幅。"""
    coord = cell.coordinate
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            total = 0.0
            for col_idx in range(rng.min_col, rng.max_col + 1):
                letter = get_column_letter(col_idx)
                dim = ws.column_dimensions.get(letter)
                width = dim.width if (dim and dim.width) else 8.43
                total += width
            return total
    dim = ws.column_dimensions.get(cell.column_letter)
    return dim.width if (dim and dim.width) else 8.43


def estimate_lines(text, width_units, chars_per_unit=0.5):
    """
    全角文字主体のテキストが何行に折り返されるかを見積もる。
    chars_per_unit: Excel列幅1単位あたりに収まる全角文字数の概算（Meiryo基準）。
    改行(\\n)は強制改行として扱う。
    """
    if not text:
        return 1
    chars_per_line = max(1, int(width_units * chars_per_unit))
    total_lines = 0
    for segment in str(text).split("\n"):
        seg_len = len(segment)
        total_lines += max(1, -(-seg_len // chars_per_line))  # ceil division
    return max(1, total_lines)


def estimate_required_height(font_size, lines, line_factor=1.45, padding=4.0):
    """1行あたりの高さ(pt) = フォントサイズ×line_factor。paddingは上下余白。"""
    return lines * font_size * line_factor + padding


def check_workbook(path, fix=False, margin=1.15, verbose=True, skip_sheets=None):
    skip_sheets = set(skip_sheets or [])
    wb = openpyxl.load_workbook(path)
    issues = []
    for sheet_name in wb.sheetnames:
        if sheet_name in skip_sheets:
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            row_idx = row[0].row
            needed_for_row = 0.0
            has_wrapped_text = False
            for cell in row:
                if cell.value is None or not isinstance(cell.value, str):
                    continue
                align = cell.alignment
                if not (align and align.wrap_text):
                    continue
                has_wrapped_text = True
                width_units = merged_width_units(ws, cell)
                font_size = (cell.font.size if cell.font and cell.font.size else 11)
                lines = estimate_lines(cell.value, width_units)
                needed = estimate_required_height(font_size, lines)
                needed_for_row = max(needed_for_row, needed)
            if not has_wrapped_text:
                continue
            actual = ws.row_dimensions[row_idx].height
            actual = actual if actual else 15.0  # openpyxl/Excel既定値
            if actual < needed_for_row * 0.9:  # 10%以上の不足のみ問題として扱う
                issues.append((sheet_name, row_idx, actual, needed_for_row))
                if fix:
                    ws.row_dimensions[row_idx].height = round(needed_for_row * margin, 1)

    if verbose:
        if issues:
            print(f"{'修正' if fix else '検出'}した行不足: {len(issues)}件")
            for sheet_name, row_idx, actual, needed in issues:
                print(f"  [{sheet_name}] row {row_idx}: 実際={actual:.1f}pt, 必要目安={needed:.1f}pt")
        else:
            print("行の高さ不足は見つかりませんでした。")

    if fix and issues:
        wb.save(path)
        print(f"保存しました: {path}")

    return issues


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", help="対象のxlsxファイルパス")
    parser.add_argument("--fix", action="store_true", help="不足している行の高さを自動修正して保存する")
    parser.add_argument("--skip-sheets", nargs="*", default=[], help="対象から除外するシート名（変更禁止の原本シート等）")
    args = parser.parse_args()
    issues = check_workbook(args.path, fix=args.fix, skip_sheets=args.skip_sheets)
    sys.exit(1 if issues and not args.fix else 0)
